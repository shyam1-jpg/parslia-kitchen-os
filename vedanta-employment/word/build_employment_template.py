#!/usr/bin/env python3
"""Build The Vedanta Kitchen Employment Word template (.docx).

Employment / kitchen operations only. No app or product content.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x1A, 0x1A, 0x1A)
GOLD = RGBColor(0x8A, 0x73, 0x4A)
DARK = RGBColor(0x22, 0x22, 0x22)
MUTED = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_BG = "1A1A1A"
GOLD_BG = "8A734A"
LIGHT_ROW = "F5F2EC"
CREAM = "F7F3EA"
WARN_BG = "F8EFE6"

HERE = Path(__file__).resolve().parent
OUT_DOCX = HERE / "Vedanta_Kitchen_Employment_Template.docx"
OUT_DOTX = HERE / "Vedanta_Kitchen_Employment_Template.dotx"


def set_run(run, *, size=11, bold=False, italic=False, color=DARK, font="Calibri", highlight=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    if highlight:
        run.font.highlight_color = highlight


def shade_cell(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def set_cell_borders(cell, color="CCCCCC"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tc_pr.append(borders)


def set_cell_margins(cell, **sides):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for side, twips in sides.items():
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(twips))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def clear_paragraph(paragraph):
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def write_cell(
    cell,
    text,
    *,
    bold=False,
    color=DARK,
    size=10,
    center=False,
    fill=None,
    highlight=None,
    font="Calibri",
):
    clear_paragraph(cell.paragraphs[0])
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, font=font, highlight=highlight)
    if fill:
        shade_cell(cell, fill)
    set_cell_borders(cell)
    set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        write_cell(table.rows[0].cells[i], header, bold=True, color=WHITE, size=10, fill=HEADER_BG)
    for r, row in enumerate(rows):
        fill = LIGHT_ROW if r % 2 else "FFFFFF"
        for c, value in enumerate(row):
            write_cell(table.rows[r + 1].cells[c], value, size=10, fill=fill)
    if col_widths:
        for row in table.rows:
            for i, width in enumerate(col_widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    if level == 1:
        set_run(run, size=16, bold=True, color=NAVY, font="Georgia")
    elif level == 2:
        set_run(run, size=13, bold=True, color=GOLD, font="Georgia")
    else:
        set_run(run, size=11, bold=True, color=NAVY)
    return p


def body(doc, text, *, bold=False, italic=False, size=11, after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=DARK)
    return p


def mixed_body(doc, parts, *, after=6):
    """parts: list of (text, kwargs)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    for text, kwargs in parts:
        run = p.add_run(text)
        set_run(run, **kwargs)
    return p


def bullet(doc, text, *, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    clear_paragraph(p)
    if bold_lead:
        run = p.add_run(bold_lead + " ")
        set_run(run, size=11, bold=True, color=DARK)
        run = p.add_run(text)
        set_run(run, size=11, color=DARK)
    else:
        run = p.add_run(text)
        set_run(run, size=11, color=DARK)
    p.paragraph_format.space_after = Pt(3)
    return p


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        clear_paragraph(p)
        run = p.add_run(item)
        set_run(run, size=11, color=DARK)
        p.paragraph_format.space_after = Pt(3)


def callout(doc, text, *, fill=CREAM, color=NAVY, bold=True):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    write_cell(cell, text, bold=bold, color=color, size=11, fill=fill)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)


def page_break(doc):
    doc.add_page_break()


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run(run, size=9, color=MUTED)


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    clear_paragraph(hp)
    run = hp.add_run("THE VEDANTA")
    set_run(run, size=9, bold=True, color=GOLD, font="Georgia")
    run = hp.add_run("   ·   Kitchen Employment Template   ·   Confidential")
    set_run(run, size=9, color=MUTED)
    hp.paragraph_format.space_after = Pt(2)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    run = fp.add_run("20-guest baseline   ·   Employment and kitchen operations only   ·   Page ")
    set_run(run, size=9, color=MUTED)
    add_page_field(fp)


def cover(doc):
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = kicker.add_run("EMPLOYMENT  ·  KITCHEN OPERATIONS  ·  WORD TEMPLATE")
    set_run(run, size=11, bold=True, color=GOLD, font="Calibri")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    run = title.add_run("The Vedanta")
    set_run(run, size=32, bold=True, color=NAVY, font="Georgia")

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Kitchen Employment Template")
    set_run(run, size=18, color=GOLD, font="Georgia")

    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("4-Week Pilot  ·  Clear Ownership  ·  One Hospitality Team")
    set_run(run, size=12, italic=True, color=MUTED, font="Georgia")

    callout(
        doc,
        "This Word file is for employment and kitchen operations only. "
        "It is not software, not a product, and not linked to any app or other project. "
        "Yellow highlighted text is a fill-in field — overtype it, then File > Save As.",
        fill=CREAM,
        color=NAVY,
    )

    add_table(
        doc,
        ["Field", "Fill in or confirm"],
        [
            ["Document title", "Kitchen Operating Response — 20-guest pilot"],
            ["To", "Shannon and Louise (Duty Managers) and the General Manager"],
            ["From", "Shyam, Head Chef"],
            ["Site / house", "The Vedanta"],
            ["Working baseline", "20 guests  (not 30–35)"],
            ["Document date", "7 September 2026"],
            ["Version", "1.0 — first issue"],
            ["Classification", "Employment / confidential"],
            ["Review date", "Week 2 of pilot, then Week 4"],
        ],
        col_widths=[5.5, 11.5],
    )

    heading(doc, "How to use this Word template", 2)
    numbered(
        doc,
        [
            "Open this file in Microsoft Word. It is A4, ready to print or Save as PDF.",
            "Yellow boxes in Part C are blank fields for a new post. Overtype them.",
            "Do not change the 20-guest baseline unless the GM reissues the house figure.",
            "Issue Part B work descriptions to each post holder and keep a signed copy (Part D).",
            "To keep a clean master: File > Save As > Word Template (*.dotx).",
        ],
    )

    heading(doc, "Contents", 2)
    for item in (
        "Part A — Professional response to the Kitchen Operating Proposal",
        "Part B — Full work descriptions at the 20-guest baseline",
        "Part C — Blank work-description template (for a new or changed post)",
        "Part D — Employment sign-off",
    ):
        bullet(doc, item)


def part_a(doc):
    page_break(doc)
    heading(doc, "Part A — Professional response")
    body(
        doc,
        "This is a professional response to The Vedanta — Kitchen Operating Proposal — "
        "Shyam, 4-Week Pilot. I support the direction. I also need the guest figure, "
        "the work descriptions and several gaps in the script corrected before we treat "
        "the paper as a working employment model.",
    )

    heading(doc, "1. Positive response", 2)
    body(
        doc,
        "Thank you for putting this on paper. The draft is the first document that treats "
        "the Kitchen as a technical section with a clear boundary, rather than a place that "
        "absorbs every hospitality gap. That is the right starting point, and I welcome it.",
    )
    body(doc, "I particularly support the following.", italic=True)
    bullet(
        doc,
        "Food quality, food safety, allergens, menus, production planning, stock and meal "
        "readiness sit with the Kitchen. Shannon and Louise can coordinate people without "
        "taking food-safety decisions they should not own.",
        bold_lead="Technical ownership is named.",
    )
    bullet(
        doc,
        "Buffet presentation, table clearing, dining-room reset, tea and coffee, guest "
        "refreshment fruit and routine restaurant crockery washing are not Kitchen work.",
        bold_lead="The Restaurant boundary is the right one.",
    )
    bullet(
        doc,
        "Developing an In-Charge and a Reliever, so the operation does not depend on one "
        "or two people, is the professional way forward.",
        bold_lead="Chef relief and succession are named.",
    )
    bullet(
        doc,
        "Head Chef work is not whatever is left after service. Menu, allergen control, "
        "ordering, training and a 7-day plan need time inside the ordinary week.",
        bold_lead="Protected management time is the right instinct.",
    )
    bullet(
        doc,
        "Kitchen can help a non-technical bottleneck when Kitchen work is complete. "
        "Neither move should put allergens, hygiene or technical standards at risk. "
        "If we cannot agree, the GM decides.",
        bold_lead="One Hospitality team, with a food-safety line, is sound.",
    )
    bullet(
        doc,
        "I do not want ghost posts. I do want enough cover to run a safe 20-guest kitchen "
        "without excessive hours.",
        bold_lead="Prove the hours before we add headcount is fair.",
    )
    bullet(
        doc,
        "A chef finishing at 21:00 must not be rostered onto the 07:00 shift the next morning. "
        "That is a working-time and safety issue, not a preference.",
        bold_lead="The close-to-open rule is correct.",
    )
    body(
        doc,
        "I want this proposal to succeed. For it to succeed as an employment document, "
        "the guest number must be the real number, each role must have a full work "
        "description, and the gaps in the script must be fixed before week one.",
    )

    heading(doc, "2. Required correction: 20 guests, not 30–35", 2)
    callout(
        doc,
        "The draft is built on a 30–35 guest baseline. The correct working baseline "
        "for this kitchen, and for this pilot, is 20 guests. That is not a small edit. "
        "It changes the model.",
        fill=WARN_BG,
        color=RGBColor(0x6B, 0x3A, 0x1F),
    )
    add_table(
        doc,
        ["Draft figure", "Correct figure", "Why it matters"],
        [
            [
                "30–35 guest pilot",
                "20 guests",
                "Prep volume, wash-up, service intensity and chef overlap all scale with covers.",
            ],
            [
                "Two chef shifts sized for 30–35",
                "Clock times may still work; do not staff as a 35-cover brigade",
                "A 4-hour overlap on 20 lunches wastes skilled hours unless used for handover and planning.",
            ],
            [
                "Three full-time KA/KP posts as the original model",
                "Confirm actual hours at 20 guests",
                "Support windows still matter. Permanent headcount must follow 20-guest demand.",
            ],
            [
                "Measures written as 30–35 guest pilot standard",
                "Measures written as 20-guest pilot standard",
                "Otherwise we judge the wrong operation and make the wrong headcount decision.",
            ],
        ],
    )
    body(
        doc,
        "Please replace 30–35 with 20 guests wherever the draft uses it. If a retreat week "
        "later rises toward 30 guests, that is a peak variant, not the baseline.",
    )

    heading(doc, "3. Full work description — Head Chef (summary)", 2)
    body(
        doc,
        "The draft lists bullets under “Shyam’s Role”. That is a purpose note, not an "
        "employment work description. The full descriptions are in Part B.",
    )
    body(
        doc,
        "Purpose. Own the Kitchen technically at a 20-guest baseline: food quality, food "
        "safety, menus, production, stock and chef standards. Lead and plan. Do not routinely "
        "carry Restaurant work, Duty Manager work, or every operational gap in the building.",
    )
    body(
        doc,
        "Reports to. General Manager for employment, performance and final dispute. Works "
        "with Shannon and Louise on live deployment. They do not override a genuine "
        "food-safety or technical Kitchen requirement.",
    )
    heading(doc, "Daily work at 20 guests", 3)
    for item in (
        "Confirm guest count, allergens, diet notes and service times for the day.",
        "Set or confirm the day’s menu, quantities and named dish owners.",
        "See that breakfast, lunch and dinner are produced to standard and handed to Restaurant on time.",
        "Check temperatures, labelling, allergen control and production hygiene.",
        "Oversee Kitchen pot wash, Kitchen cleaning and the technical close.",
        "Review shortages and the order list; authorise or place Kitchen orders as required.",
        "Brief the chef on duty and the KA/KP support; take a proper handover.",
        "Record any unplanned extra hour or support move, with the reason.",
        "Complete the Head Chef planning block when the rota has protected it.",
    ):
        bullet(doc, item)
    heading(doc, "What this role does not own", 3)
    for item in (
        "Restaurant buffet setup, guest-facing presentation, table clearing or dining-room reset.",
        "Tea, coffee, refreshment area, guest biscuits, refreshment fruit or FOH milk.",
        "Routine Restaurant crockery and mug washing.",
        "Reception, check-in, housekeeping or Estates work.",
        "The Duty Manager function. Shannon and Louise hold that.",
    ):
        bullet(doc, item)

    heading(doc, "4. Issues found in the script", 2)
    body(
        doc,
        "The script is well intended. It is not yet safe to issue as an employment or rota document.",
    )

    issues = [
        (
            "4.1 The guest number is wrong",
            "The purpose, chef-shift heading and measures table all say 30–35 guests. "
            "This kitchen’s working baseline is 20 guests. If we pilot the 35-cover model "
            "on a 20-guest house, we will either waste skilled hours or conclude we need "
            "people we do not need. Fix: rewrite the paper for 20 guests. Treat any week "
            "near 30 as a recorded peak.",
        ),
        (
            "4.2 “Do not carry every gap” is then contradicted",
            "Section 1 says I should lead and plan, then gives me quality, safety, menus, "
            "production, stock, waste, food cost, both deputy posts and the staffing "
            "requirement. That is the correct Head Chef scope only if Chef 2 and the Reliever "
            "have real authority and management time is on the rota. Fix: name the deputy "
            "on every shift.",
        ),
        (
            "4.3 Three chef names, none tied to the shift table",
            "Head Chef, In-Charge / Chef 2, Reliever, Morning Chef and Late Chef are mixed. "
            "Posts are people. Morning and Late are shifts. A rota clerk cannot build a week "
            "from this.",
        ),
        (
            "4.4 The Reliever is required, but is not a real post yet",
            "The paper needs cover for off-days, then asks after four weeks whether the "
            "Reliever is internal or external. If the Reliever is only a hope for week 5, "
            "weeks 1–4 fail by design. Fix: name the Reliever before day one.",
        ),
        (
            "4.5 KA/KP headcount is three ideas at once",
            "Three full-time posts, “confirm actual hours”, and a float principle sit in one "
            "paragraph. Kitchen Assistant and Kitchen Porter are different skills and are "
            "never defined. There is no 7-day worked example.",
        ),
        (
            "4.6 Four-hour overlap is a 30–35 idea",
            "At 20 guests, two chefs both producing lunch idle skilled time. Keep the clock "
            "times if they fit service. One chef owns lunch; the other uses 12:00–16:00 for "
            "handover, quality, next-meal prep and Head Chef planning.",
        ),
        (
            "4.7 Planning time is hoped for, not rostered",
            "“Where practical” and a weekly admin day that “may be used” is how planning "
            "disappears. Roster a named block with a template: 7-day plan, 48-hour "
            "confirmation, order list, support requirement.",
        ),
        (
            "4.8 Ownership leaks will put Restaurant work back on Kitchen",
            "Same fruit / same knife; “help when directed” for FOH wash; nobody named to "
            "declare Kitchen work complete; no owner for a shared dishwasher; FOH milk "
            "versus Kitchen dairy. The leading chef must confirm Kitchen work complete "
            "against a close checklist.",
        ),
        (
            "4.9 Duty Manager versus Head Chef is too loose",
            "“Genuine” food safety is undefined. There is no live RACI for the pass, a late "
            "service or an 86’d dish. “No separate Duty Manager role” reads as if Duty "
            "Manager work has been removed. Shannon and Louise are the Duty Managers; "
            "no third DM post is proposed. Until the GM decides, the food-safety position stands.",
        ),
        (
            "4.10 Ordering is a principle, not a system",
            "The list, authorised orderers, pars, cut-offs and waste record are unnamed. "
            "Name at least two authorised Kitchen orderers plus the Head Chef, and train "
            "them before the pilot.",
        ),
        (
            "4.11 Pilot measures are not measurable",
            "Service times are missing. “Agreed quality”, “excessive hours” and “routinely” "
            "have no definition. There is no scorecard, no week-2 review, and no pass mark "
            "for week 4.",
        ),
        (
            "4.12 The decision section has no pass mark",
            "Adopt / adjust / confirm Reliever / confirm KA mix is a process, not a standard. "
            "Agree numbers before day one.",
        ),
        (
            "4.13 Employment content is missing",
            "Service times, diet path, close checklist, a worked 20-guest rota, breaks, "
            "Working Time rest (21:00 to 07:00 is 10 hours, not 11), overtime or TOIL, "
            "competency sign-off, and a peak plan if the house jumps toward 30.",
        ),
        (
            "4.14 Language that will cause arguments later",
            "KA/KP is used as one job. The “original pilot model” is mentioned and not "
            "attached. Morning 07:00 and late 21:00 cannot be judged without service times.",
        ),
    ]
    for title, text in issues:
        heading(doc, title, 3)
        body(doc, text)

    heading(doc, "5. What I recommend we issue instead", 2)
    numbered(
        doc,
        [
            "Reissue the proposal at 20 guests. Delete 30–35 from purpose, shift titles and the measures table.",
            "Attach the full work descriptions in Part B as the employment attachment.",
            "Name the Reliever and the authorised orderers before week 1.",
            "Roster the overlap as handover and Head Chef planning, not as two chefs producing 20 lunches.",
            "Roster support as two windows plus relief, and let the 20-guest hours decide whether three full-time KA/KP posts are justified.",
            "Issue a one-page scorecard and a week-2 review, then decide at week 4 against agreed numbers.",
            "Keep the principle that is already right: Kitchen owns the food and the technical Kitchen standard. Restaurant owns guest-facing dining and refreshment. I define the Kitchen requirement. Shannon and Louise coordinate live deployment. The GM decides a deadlock.",
        ],
    )
    body(doc, "I am ready to run the 4-week pilot on that basis.", italic=True)
    body(doc, "Shyam")
    body(doc, "Head Chef", after=0)
    body(doc, "The Vedanta")


def role_block(doc, title, fields, purpose, daily, weekly=None, not_do=None, success=None):
    heading(doc, title, 2)
    add_table(doc, ["Field", "Detail"], fields, col_widths=[5.5, 11.5])
    heading(doc, "Purpose of the role", 3)
    body(doc, purpose)
    heading(doc, "Daily duties", 3)
    numbered(doc, daily)
    if weekly:
        heading(doc, "Weekly and periodic duties", 3)
        numbered(doc, weekly)
    if not_do:
        heading(doc, "What this role does not do", 3)
        for item in not_do:
            bullet(doc, item)
    if success:
        heading(doc, "Success in the role", 3)
        if isinstance(success, list):
            for item in success:
                bullet(doc, item)
        else:
            body(doc, success)


def part_b(doc):
    page_break(doc)
    heading(doc, "Part B — Full work descriptions (20 guests)")
    body(
        doc,
        "These descriptions replace the short bullet lists in the draft proposal. "
        "They are for employment, rota, induction and the 4-week pilot.",
    )
    heading(doc, "Common rules for every Kitchen post", 2)
    for item in (
        "The working house is 20 guests. A higher retreat week is a recorded peak, not the normal job size.",
        "Kitchen owns food production, food safety, Kitchen stock, Kitchen pot wash and Kitchen cleaning.",
        "Restaurant owns guest-facing dining, buffet presentation, tea and coffee, refreshment fruit and routine FOH wash-up.",
        "No cross-deployment may put allergens, hygiene or technical standards at risk.",
        "A chef who finishes at 21:00 is not rostered to start at 07:00 the next day.",
        "Extra hours and temporary support moves are recorded with a reason on the day.",
    ):
        bullet(doc, item)

    role_block(
        doc,
        "1. Head Chef — Shyam",
        [
            ["Job title", "Head Chef"],
            ["Post holder", "Shyam"],
            ["Department", "Kitchen"],
            ["Reports to", "General Manager"],
            ["Works with", "Kitchen In-Charge, Reliever Chef, KA/KP, Duty Managers, Restaurant lead"],
            ["Responsible for", "Technical Kitchen standard, quality, safety, menus, production, stock, training, staffing requirement"],
            ["Contract pattern (pilot)", "Rostered chef shifts plus a protected planning block inside the ordinary week"],
            ["Baseline", "20 guests"],
        ],
        "Lead the Kitchen as a technical section of Hospitality. Plan the food and the Kitchen "
        "team for 20 guests so breakfast, lunch and dinner are ready, safe and consistent. "
        "Build the In-Charge and Reliever so the Kitchen does not depend on one person. "
        "Do not absorb Restaurant, reception, housekeeping or Duty Manager work.",
        [
            "Read the day picture: 20-guest count, arrivals, allergens, diet notes, service times.",
            "Confirm the menu, production quantities and named dish owners.",
            "Brief the Kitchen team at the start of the relevant shift.",
            "Own or oversee production of the meals on the shifts you are rostered to lead.",
            "Check quality, presentation, temperatures, labelling and allergen control before handover.",
            "Hand meals to Restaurant at the agreed time. Kitchen work stops at a ready, labelled, safe handover.",
            "Review the shortage / order list. Authorise or place Kitchen orders due that day.",
            "See that Kitchen pot wash, surfaces and the production close are done to standard.",
            "Take and give a written or board handover.",
            "Report equipment or building faults to Estates. Do not own the repair.",
            "Record any extra hour or support movement and the reason.",
            "Use the protected planning block, when rostered, for the 7-day plan, 48-hour confirmation, orders and the support requirement.",
        ],
        weekly=[
            "Issue a 7-day plan: menu, 20-guest quantities, prep owners, stock and orders, chef and KA/KP allocation, allergens, risks, support requirement.",
            "Issue a 48-hour confirmation: final count, substitutions, deliveries, service timing, any targeted support change.",
            "Keep a 4-week outlook as visibility only.",
            "Train and assess the In-Charge and Reliever against the same technical standard.",
            "Review waste, food-cost movement and recurring bottlenecks with the GM as required.",
            "Give Shannon and Louise the Kitchen staffing and support requirement in time for the rota.",
            "Review competency sign-off for new or relief KA/KP staff before they work unsupervised.",
        ],
        not_do=[
            "Restaurant buffet setup, table clearing, dining-room reset or guest refreshment service.",
            "Stocking tea, coffee, FOH milk, sugar, biscuits or refreshment fruit.",
            "Routine Restaurant crockery and mug washing.",
            "Reception, guest administration, housekeeping or Estates maintenance.",
            "Acting as Duty Manager for the wider building.",
        ],
        success=[
            "Breakfast, lunch and dinner handed over on time.",
            "Agreed quality and presentation held.",
            "No safety or allergen failure caused by labour or shortcut.",
            "Weekly off-days held; extra hours recorded and not routine.",
            "7-day and 48-hour plans issued in time for staffing and ordering.",
            "In-Charge and Reliever can lead a shift without the Head Chef standing behind every dish.",
        ],
    )

    role_block(
        doc,
        "2. Kitchen In-Charge / Chef 2",
        [
            ["Job title", "Kitchen In-Charge (Chef 2)"],
            ["Department", "Kitchen"],
            ["Reports to", "Head Chef"],
            ["Responsible for", "Day-to-day production lead on the shifts they are rostered to run"],
            ["Baseline", "20 guests"],
        ],
        "Run the operational floor when rostered as the leading chef. Deliver the Head Chef’s "
        "menu and plan to the same standard. This is a deputy post with real authority on the "
        "shift, not a helper who waits for Shyam on every decision.",
        [
            "Take the handover and the day’s 20-guest plan, allergens and dish list.",
            "Run mise en place for the meals on that shift.",
            "Cook and finish dishes to the agreed recipes and presentation.",
            "Check temperatures, labels, allergen separation and hygiene during production.",
            "Have food ready for the agreed handover time and pass it to Restaurant.",
            "Direct the KA/KP on Kitchen work only, unless a recorded one-off move is agreed and Kitchen work is complete.",
            "Complete Kitchen close tasks for that shift.",
            "Write the handover: what is left, what is short, what the next chef must know.",
            "Escalate genuine food-safety or staffing risk without standing down the safety position.",
            "When the Head Chef is off, authorise routine Kitchen decisions inside the 7-day plan.",
        ],
        not_do=[
            "Redesign the menu against the Head Chef’s plan, except for a recorded safety or shortage substitution.",
            "Own Restaurant service or refreshment.",
            "Override the close-to-open rest rule to “help”.",
        ],
        success="The shift runs to time and standard when the Head Chef is not on the floor. Handover is complete. KA/KP are used for Kitchen support, not as a default FOH crew.",
    )

    role_block(
        doc,
        "3. Skilled Reliever Chef",
        [
            ["Job title", "Skilled Reliever Chef"],
            ["Department", "Kitchen"],
            ["Reports to", "Head Chef"],
            ["Responsible for", "Cover of either main chef shift to the same technical standard"],
            ["Baseline", "20 guests"],
        ],
        "Give the Kitchen resilience. Cover Head Chef or In-Charge off-days, short absence and "
        "leave. The Reliever must be able to run either 07:00–16:00 or 12:00–21:00 without a "
        "drop in safety, quality or timing. This post must be named and competent before the "
        "4-week pilot.",
        [
            "When covering: carry out the full leading-chef work description for that shift.",
            "Follow the issued 7-day plan and 48-hour confirmation.",
            "Keep recipes, allergen control and handover at the same standard as a normal day.",
            "Do not leave unfinished production, unlabelled food or an incomplete close because “I am only covering”.",
            "Record any substitution, shortage or extra hour.",
            "When not covering: work designated services to sign off both shift patterns, learn the order routine, and complete competency sign-off.",
        ],
        not_do=[
            "Exist only on paper as a week-5 decision while the Head Chef covers every off-day in weeks 1–4.",
            "Be used as extra KA/KP labour on a 20-guest day when a chef is already on the shift.",
        ],
        success="Off-days and short absence are covered without excessive hours on the remaining chef. Quality and safety do not dip on Reliever-led shifts.",
    )

    role_block(
        doc,
        "4. Kitchen Assistant / Kitchen Porter — Early support",
        [
            ["Job title", "Kitchen Assistant / Kitchen Porter — Early"],
            ["Typical window", "07:00–15:00"],
            ["Department", "Kitchen"],
            ["Reports to", "Chef leading the morning shift"],
            ["Responsible for", "Morning Kitchen support for a 20-guest house"],
            ["Baseline", "20 guests"],
        ],
        "Support breakfast and lunch production so the chef’s skilled time is used on food, "
        "not on porter work. KA work is prep assistance under the chef. KP work is pot wash, "
        "bins, floors and scheduled cleaning. They are not the same job. The rota must say "
        "which skill is required that day.",
        [
            "Clock in, uniform, hand wash, read the day board.",
            "Assist morning prep for breakfast and lunch at 20-guest quantities, not a 35-cover batch.",
            "Receive deliveries with the chef or authorised person: check, date, put away, FIFO.",
            "Move stock, label containers, keep walkways clear.",
            "Wash pots, pans and Kitchen production equipment through the morning and after lunch.",
            "Clean Kitchen surfaces, sinks and allocated areas. Follow clean-as-you-go.",
            "Do not start Restaurant buffet dressing, tea-station fruit or FOH mug washing as a routine task.",
            "If Kitchen work is complete and the leading chef confirms it, you may be moved to an agreed non-technical bottleneck. That move is recorded.",
        ],
        success="Chef is not pulled onto repetitive wash-up or delivery runs during breakfast and lunch. Kitchen is clean, labelled and ready for the late window handover.",
    )

    role_block(
        doc,
        "5. Kitchen Assistant / Kitchen Porter — Late support",
        [
            ["Job title", "Kitchen Assistant / Kitchen Porter — Late"],
            ["Typical window", "15:00–22:00"],
            ["Department", "Kitchen"],
            ["Reports to", "Chef leading the late / dinner shift"],
            ["Responsible for", "Dinner support, Kitchen pot wash and Kitchen close"],
            ["Baseline", "20 guests"],
        ],
        "Support dinner production and close the Kitchen to a safe standard so the late chef "
        "is not left with porter work after service.",
        [
            "Take handover from the early support and the morning chef.",
            "Assist dinner prep as directed at 20-guest quantities.",
            "Run Kitchen pot wash through dinner and close.",
            "Complete the Kitchen close list: equipment, surfaces, floors, bins, labels, fridges as required.",
            "Do not take on dining-room clear-down or Restaurant crockery as a routine finish.",
            "Leave a clean, safe kitchen for the 07:00 chef. Do not leave production work that should have been finished at 21:00.",
        ],
        success="Dinner support is real, not cosmetic. Kitchen close is finished inside the planned window. The morning chef does not inherit last night’s pots.",
    )

    role_block(
        doc,
        "6. Kitchen Assistant / Kitchen Porter — Relief / Float",
        [
            ["Job title", "Kitchen Assistant / Kitchen Porter — Relief / Float"],
            ["Typical use", "Days off, sickness, annual leave, named peak bottlenecks"],
            ["Department", "Kitchen"],
            ["Reports to", "Head Chef for allocation; chef leading the shift when on duty"],
            ["Baseline", "20 guests"],
        ],
        "Cover the early or late window when the regular support is off, and attend a named "
        "peak. This is not a third person standing in the kitchen every day “in case”.",
        [
            "Cover the early or late work description in full.",
            "Or complete a named task list: deep clean, delivery day, wash-up spike, stock day.",
            "Do not invent extra work on a quiet 20-guest day to fill the shift.",
        ],
        success="Support windows remain covered when someone is off. Extra hours on the regular KA/KP and chefs do not become the default cover method.",
    )

    heading(doc, "7. How the posts sit on a 20-guest day", 2)
    body(doc, "These are shifts, not extra job titles.")
    add_table(
        doc,
        ["Shift", "Clock", "Who may hold it", "Rule at 20 guests"],
        [
            ["Morning chef", "07:00–16:00", "Head Chef, In-Charge or Reliever", "Breakfast, lunch production, readiness, handover."],
            ["Late / dinner chef", "12:00–21:00", "Head Chef, In-Charge or Reliever", "Dinner, technical close, handover. Not rostered to 07:00 the next day."],
            ["Overlap", "12:00–16:00", "Both chefs on duty that day", "One chef owns lunch. The other does handover, quality, next-meal prep and Head Chef planning."],
            ["Early support", "07:00–15:00", "KA and/or KP as rostered", "Kitchen support only unless a recorded one-off is agreed."],
            ["Late support", "15:00–22:00", "KA and/or KP as rostered", "Dinner support and Kitchen close."],
            ["Relief / Float", "As required", "Relief post", "Cover a window or a named peak only — not a third body every day."],
        ],
    )

    heading(doc, "8. Duty Manager interface (not a Kitchen post)", 2)
    body(
        doc,
        "Shannon and Louise are the Duty Managers for the wider Hospitality operation. "
        "This is not a Kitchen job and no extra Duty Manager post is proposed.",
    )
    add_table(
        doc,
        ["They do", "They do not"],
        [
            ["Review the whole house and coordinate live deployment", "Override a genuine Kitchen food-safety or technical requirement"],
            ["Receive the Head Chef’s weekly support requirement", "Redesign the menu, the 20-guest production plan or the Kitchen close standard"],
            ["Record a one-off KA/KP move when Kitchen work is complete", "Treat Kitchen as the default Restaurant wash-up crew"],
            ["Escalate a deadlock to the GM", "Leave a shift without a competent chef in order to cover FOH"],
        ],
    )
    body(
        doc,
        "Until the GM decides a deadlock, the food-safety position set by the chef leading the shift stands.",
    )


def highlight_row(table, row_index, value_col=1, placeholder="[Type here]"):
    cell = table.rows[row_index].cells[value_col]
    write_cell(
        cell,
        placeholder,
        size=10,
        fill="FFF3B0",
        highlight=WD_COLOR_INDEX.YELLOW,
    )


def part_c(doc):
    page_break(doc)
    heading(doc, "Part C — Blank work-description template")
    body(
        doc,
        "Use this page when a new Kitchen post is created or an existing post changes. "
        "Overtype every yellow field. Keep the 20-guest baseline unless the GM reissues it. "
        "Copy this page in Word (right-click the page break, or copy the tables) for a second post.",
    )
    callout(
        doc,
        "Yellow = fill in. Do not leave yellow text in a signed employment copy.",
        fill="FFF3B0",
        color=NAVY,
    )

    table = add_table(
        doc,
        ["Field", "Detail"],
        [
            ["Job title", "[Job title]"],
            ["Post holder", "[Name, or Vacant]"],
            ["Department", "Kitchen"],
            ["Reports to", "[Head Chef / leading chef]"],
            ["Works with", "[List]"],
            ["Responsible for", "[One sentence]"],
            ["Typical hours / window", "[e.g. 07:00–16:00]"],
            ["Contract / pattern", "[Full-time / part-time / relief]"],
            ["Working baseline", "20 guests"],
            ["Start date", "[Date]"],
            ["Review date", "[Week 2 / Week 4 / probation]"],
        ],
        col_widths=[5.5, 11.5],
    )
    for i in (1, 2, 4, 5, 6, 7, 8, 10, 11):
        highlight_row(table, i)

    heading(doc, "Purpose of the role", 3)
    p = doc.add_paragraph()
    run = p.add_run("[Write the purpose in 3–5 lines. Say what this post owns at 20 guests, and what it must not absorb from Restaurant or Duty Manager.]")
    set_run(run, size=11, color=DARK, highlight=WD_COLOR_INDEX.YELLOW)

    heading(doc, "Daily duties", 3)
    for _ in range(8):
        p = doc.add_paragraph(style="List Number")
        clear_paragraph(p)
        run = p.add_run("[Duty]")
        set_run(run, size=11, color=DARK, highlight=WD_COLOR_INDEX.YELLOW)
        p.paragraph_format.space_after = Pt(3)

    heading(doc, "Weekly / periodic duties", 3)
    for _ in range(4):
        p = doc.add_paragraph(style="List Number")
        clear_paragraph(p)
        run = p.add_run("[Duty]")
        set_run(run, size=11, color=DARK, highlight=WD_COLOR_INDEX.YELLOW)
        p.paragraph_format.space_after = Pt(3)

    heading(doc, "What this role does not do", 3)
    for _ in range(4):
        p = doc.add_paragraph(style="List Bullet")
        clear_paragraph(p)
        run = p.add_run("[Out of scope]")
        set_run(run, size=11, color=DARK, highlight=WD_COLOR_INDEX.YELLOW)
        p.paragraph_format.space_after = Pt(3)

    heading(doc, "Success in the role (20-guest pilot)", 3)
    p = doc.add_paragraph()
    run = p.add_run("[How this post will be judged: on-time handover, safety, quality, recorded extra hours, cover of off-days.]")
    set_run(run, size=11, color=DARK, highlight=WD_COLOR_INDEX.YELLOW)

    heading(doc, "Authorised to (tick / complete)", 3)
    add_table(
        doc,
        ["Authority", "Yes / No", "Notes"],
        [
            ["Lead a chef shift unsupervised", "[Yes/No]", "[ ]"],
            ["Authorise a Kitchen order", "[Yes/No]", "[ ]"],
            ["Confirm Kitchen work complete", "[Yes/No]", "[ ]"],
            ["Hold or 86 a dish inside the 7-day plan", "[Yes/No]", "[ ]"],
            ["Direct KA/KP on the shift", "[Yes/No]", "[ ]"],
        ],
    )
    auth = doc.tables[-1]
    for r in range(1, 6):
        write_cell(auth.rows[r].cells[1], "[Yes/No]", size=10, fill="FFF3B0", highlight=WD_COLOR_INDEX.YELLOW)
        write_cell(auth.rows[r].cells[2], "[ ]", size=10, fill="FFF3B0", highlight=WD_COLOR_INDEX.YELLOW)


def part_d(doc):
    page_break(doc)
    heading(doc, "Part D — Employment sign-off")
    body(
        doc,
        "I have read this Kitchen Employment Template at the 20-guest baseline. "
        "I understand the work description for my post, what Kitchen owns, and what "
        "Kitchen does not own. I understand that a chef finishing at 21:00 is not "
        "rostered to start at 07:00 the next day.",
    )
    add_table(
        doc,
        ["", "Post holder", "Head Chef", "General Manager"],
        [
            ["Name", "", "Shyam", ""],
            ["Job title", "", "Head Chef", "General Manager"],
            ["Signature", "", "", ""],
            ["Date", "", "", ""],
        ],
    )
    sign = doc.tables[-1]
    for r in range(1, 5):
        for c in (1, 3):
            write_cell(sign.rows[r].cells[c], "", size=10, fill="FFF3B0")
        if r in (3, 4):
            write_cell(sign.rows[r].cells[2], "", size=10, fill="FFF3B0")

    heading(doc, "Simple rule", 2)
    callout(
        doc,
        "Kitchen owns the food and the technical Kitchen standard. "
        "Restaurant owns the guest-facing dining and refreshment service. "
        "Shyam defines the Kitchen requirement. Shannon and Louise coordinate live "
        "deployment as Duty Managers. The GM decides a deadlock.",
        fill=CREAM,
        color=NAVY,
    )


def set_core_properties(doc):
    props = doc.core_properties
    props.title = "The Vedanta — Kitchen Employment Template (20 guests)"
    props.subject = "Employment and kitchen operations — not a software or product document"
    props.author = "Shyam, Head Chef"
    props.category = "Employment"
    props.comments = (
        "Standalone Kitchen employment Word template for The Vedanta. "
        "Working baseline 20 guests, not 30-35. No app or product content."
    )


def write_dotx_copy(docx_path: Path, dotx_path: Path):
    """Save a Word template copy (.dotx) with the correct content type."""
    import zipfile
    from io import BytesIO

    buf = BytesIO()
    with zipfile.ZipFile(docx_path, "r") as src, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                text = data.decode("utf-8")
                text = text.replace(
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                )
                data = text.encode("utf-8")
            dst.writestr(item, data)
    dotx_path.write_bytes(buf.getvalue())


def main():
    doc = Document()
    setup_document(doc)
    set_core_properties(doc)
    cover(doc)
    part_a(doc)
    part_b(doc)
    part_c(doc)
    part_d(doc)
    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    write_dotx_copy(OUT_DOCX, OUT_DOTX)
    print(f"Wrote {OUT_DOCX}")
    print(f"Wrote {OUT_DOTX}")


if __name__ == "__main__":
    main()
